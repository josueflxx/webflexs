"""Professional Excel exports for immutable fiscal documents."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


ORANGE = "FF6B35"
INK = "172033"
INK_SOFT = "445066"
PAPER = "FFFFFF"
SURFACE = "F3F6FA"
SURFACE_ALT = "E8EDF4"
LINE = "CBD5E1"

MONEY_FORMAT = '$ #,##0.00;[Red]-$ #,##0.00'
QUANTITY_FORMAT = "0.###"
RATE_FORMAT = "0.00%"
DATE_FORMAT = "dd/mm/yyyy"


def _decimal(value, default="0"):
    try:
        return Decimal(str(value if value not in (None, "") else default))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal(default)


def _snapshot_section(snapshot, key):
    value = snapshot.get(key, {}) if isinstance(snapshot, dict) else {}
    return value if isinstance(value, dict) else {}


def _safe_text(value, fallback="-"):
    text = str(value or "").strip()
    return text or fallback


def _date_value(value):
    if not value:
        return None
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        try:
            return date.fromisoformat(text)
        except ValueError:
            return text


def _document_number(document):
    point = _safe_text(
        getattr(document, "point_of_sale_number_snapshot", "")
        or getattr(getattr(document, "point_of_sale", None), "number", ""),
        "",
    )
    if getattr(document, "number", None) is not None:
        if point:
            return f"{point.zfill(5)}-{str(document.number).zfill(8)}"
        return str(document.number).zfill(8)
    return _safe_text(getattr(document, "external_number", ""))


def build_fiscal_export_filename(document):
    doc_label = slugify(getattr(document, "doc_type", "comprobante")) or "comprobante"
    number = slugify(_document_number(document)) or str(document.pk)
    return f"{doc_label}-{number}.xlsx"


def build_fiscal_workbook(document, *, snapshot=None):
    """Build a read-only-friendly workbook from the frozen fiscal snapshot."""

    snapshot = snapshot if isinstance(snapshot, dict) else {}
    emitter = _snapshot_section(snapshot, "emitter")
    client = _snapshot_section(snapshot, "client")
    operation = _snapshot_section(snapshot, "operation")
    totals = _snapshot_section(snapshot, "totals")

    wb = Workbook()
    ws = wb.active
    ws.title = "Factura"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A16"
    ws.auto_filter.ref = None

    thin = Side(style="thin", color=LINE)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=INK)
    orange_fill = PatternFill("solid", fgColor=ORANGE)
    soft_fill = PatternFill("solid", fgColor=SURFACE)
    alt_fill = PatternFill("solid", fgColor=SURFACE_ALT)

    for col, width in {
        "A": 15,
        "B": 13,
        "C": 42,
        "D": 17,
        "E": 15,
        "F": 14,
        "G": 15,
        "H": 18,
    }.items():
        ws.column_dimensions[col].width = width

    company_name = _safe_text(
        emitter.get("legal_name")
        or emitter.get("name")
        or getattr(getattr(document, "company", None), "name", "")
    )
    document_label = _safe_text(document.get_doc_type_display())
    number_display = _document_number(document)

    ws.merge_cells("A1:H1")
    ws["A1"] = company_name
    ws["A1"].font = Font(color=PAPER, bold=True, size=18)
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:H2")
    ws["A2"] = f"{document_label}  {number_display}"
    ws["A2"].font = Font(color=PAPER, bold=True, size=14)
    ws["A2"].fill = orange_fill
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 26

    ws.merge_cells("A3:H3")
    ws["A3"] = (
        f"Estado: {_safe_text(document.get_status_display())}  |  "
        f"Moneda: {_safe_text(getattr(document, 'currency', 'ARS'), 'ARS')}  |  "
        "Exportación informativa del comprobante registrado"
    )
    ws["A3"].font = Font(color=INK_SOFT, italic=True, size=10)
    ws["A3"].fill = soft_fill
    ws["A3"].alignment = Alignment(vertical="center")

    info_rows = [
        (
            "Emisor",
            company_name,
            "Cliente",
            _safe_text(client.get("name")),
        ),
        (
            "CUIT emisor",
            _safe_text(emitter.get("cuit") or getattr(document, "issuer_cuit_snapshot", "")),
            _safe_text(client.get("document_type_label"), "CUIT/DNI"),
            _safe_text(client.get("document_number")),
        ),
        (
            "Condición fiscal",
            _safe_text(emitter.get("tax_condition_label")),
            "Condición de IVA",
            _safe_text(client.get("tax_condition_label")),
        ),
        (
            "Domicilio fiscal",
            " / ".join(
                filter(
                    None,
                    [
                        str(emitter.get("fiscal_address") or "").strip(),
                        str(emitter.get("fiscal_city") or "").strip(),
                        str(emitter.get("fiscal_province") or "").strip(),
                    ],
                )
            )
            or "-",
            "Domicilio cliente",
            " / ".join(
                filter(
                    None,
                    [
                        str(client.get("fiscal_address") or "").strip(),
                        str(client.get("fiscal_city") or "").strip(),
                        str(client.get("fiscal_province") or "").strip(),
                    ],
                )
            )
            or "-",
        ),
        (
            "Fecha de emisión",
            _date_value(getattr(document, "issued_at", None)),
            "Condición de venta",
            "Cuenta corriente" if operation.get("billing_mode") == "official" else "-",
        ),
        (
            "CAE",
            _safe_text(getattr(document, "cae", "")),
            "Vencimiento CAE",
            _date_value(getattr(document, "cae_due_date", None)),
        ),
    ]

    start_row = 5
    for row_index, (left_label, left_value, right_label, right_value) in enumerate(
        info_rows, start=start_row
    ):
        ws.cell(row=row_index, column=1, value=left_label)
        ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=4)
        ws.cell(row=row_index, column=2, value=left_value)
        ws.cell(row=row_index, column=5, value=right_label)
        ws.merge_cells(start_row=row_index, start_column=6, end_row=row_index, end_column=8)
        ws.cell(row=row_index, column=6, value=right_value)
        for col in range(1, 9):
            cell = ws.cell(row=row_index, column=col)
            cell.fill = soft_fill if row_index % 2 else alt_fill
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=row_index, column=1).font = Font(color=INK_SOFT, bold=True)
        ws.cell(row=row_index, column=5).font = Font(color=INK_SOFT, bold=True)
        for value_cell in (ws.cell(row=row_index, column=2), ws.cell(row=row_index, column=6)):
            value_cell.font = Font(color=INK, bold=True)
            if isinstance(value_cell.value, (date, datetime)):
                value_cell.number_format = DATE_FORMAT
        ws.row_dimensions[row_index].height = 24

    item_header_row = 14
    headers = [
        "Código / SKU",
        "Cantidad",
        "Producto",
        "Precio neto",
        "Descuento",
        "IVA",
        "Importe IVA",
        "Importe total",
    ]
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=item_header_row, column=col_index, value=header)
        cell.fill = header_fill
        cell.font = Font(color=PAPER, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[item_header_row].height = 30

    items = list(document.items.all().order_by("line_number"))
    first_item_row = item_header_row + 1
    current_row = first_item_row
    if items:
        for index, item in enumerate(items):
            values = [
                _safe_text(item.sku),
                float(_decimal(item.quantity)),
                _safe_text(item.description),
                float(_decimal(item.unit_price_net)),
                float(_decimal(item.discount_percentage) / Decimal("100")),
                float(_decimal(item.iva_rate) / Decimal("100")),
                float(_decimal(item.iva_amount)),
                float(_decimal(item.total_amount)),
            ]
            for col_index, value in enumerate(values, start=1):
                cell = ws.cell(row=current_row, column=col_index, value=value)
                cell.fill = PatternFill("solid", fgColor=PAPER if index % 2 == 0 else SURFACE)
                cell.border = border
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=col_index in {1, 3},
                    horizontal="right" if col_index in {2, 4, 5, 6, 7, 8} else "left",
                )
            ws.cell(current_row, 2).number_format = QUANTITY_FORMAT
            ws.cell(current_row, 4).number_format = MONEY_FORMAT
            ws.cell(current_row, 5).number_format = RATE_FORMAT
            ws.cell(current_row, 6).number_format = RATE_FORMAT
            ws.cell(current_row, 7).number_format = MONEY_FORMAT
            ws.cell(current_row, 8).number_format = MONEY_FORMAT
            ws.row_dimensions[current_row].height = 30
            current_row += 1
        ws.auto_filter.ref = f"A{item_header_row}:H{current_row - 1}"
    else:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
        cell = ws.cell(row=current_row, column=1, value="El comprobante no tiene ítems fiscales registrados.")
        cell.fill = soft_fill
        cell.font = Font(color=INK_SOFT, italic=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.row_dimensions[current_row].height = 30
        current_row += 1

    totals_start = current_row + 2
    subtotal = _decimal(totals.get("subtotal_net"), document.subtotal_net)
    discount = _decimal(totals.get("discount_total"), document.discount_total)
    tax = _decimal(totals.get("tax_total"), document.tax_total)
    total = _decimal(totals.get("total"), document.total)
    net = max(subtotal - discount, Decimal("0"))
    summary = [
        ("Subtotal neto", subtotal),
        ("Descuento", discount),
        ("Neto gravado", net),
        ("IVA", tax),
        ("TOTAL", total),
    ]
    for offset, (label, value) in enumerate(summary):
        row = totals_start + offset
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
        label_cell = ws.cell(row=row, column=6, value=label)
        value_cell = ws.cell(row=row, column=8, value=float(value))
        label_cell.font = Font(color=PAPER if label == "TOTAL" else INK, bold=True)
        value_cell.font = Font(color=PAPER if label == "TOTAL" else INK, bold=True, size=12 if label == "TOTAL" else 10)
        label_cell.alignment = Alignment(horizontal="right", vertical="center")
        value_cell.alignment = Alignment(horizontal="right", vertical="center")
        value_cell.number_format = MONEY_FORMAT
        for col in range(6, 9):
            ws.cell(row=row, column=col).fill = orange_fill if label == "TOTAL" else soft_fill
            ws.cell(row=row, column=col).border = border
        ws.row_dimensions[row].height = 25 if label == "TOTAL" else 21

    observations = "\n".join(
        bit
        for bit in [
            str(operation.get("notes") or "").strip(),
            str(operation.get("admin_notes") or "").strip(),
        ]
        if bit
    )
    notes_row = totals_start
    ws.merge_cells(start_row=notes_row, start_column=1, end_row=notes_row, end_column=5)
    ws.cell(notes_row, 1, "Observaciones")
    ws.cell(notes_row, 1).fill = header_fill
    ws.cell(notes_row, 1).font = Font(color=PAPER, bold=True)
    ws.cell(notes_row, 1).alignment = Alignment(vertical="center")
    for col in range(1, 6):
        ws.cell(notes_row, col).fill = header_fill
        ws.cell(notes_row, col).border = border
    ws.merge_cells(start_row=notes_row + 1, start_column=1, end_row=notes_row + 4, end_column=5)
    notes_cell = ws.cell(notes_row + 1, 1, observations or "Sin observaciones.")
    notes_cell.fill = soft_fill
    notes_cell.font = Font(color=INK_SOFT)
    notes_cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in range(notes_row + 1, notes_row + 5):
        for col in range(1, 6):
            ws.cell(row, col).fill = soft_fill
            ws.cell(row, col).border = border

    footer_row = max(notes_row + 6, totals_start + len(summary) + 1)
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=8)
    ws.cell(
        footer_row,
        1,
        "Documento exportado desde FLEXS. Los datos fiscales provienen del comprobante almacenado y no son editados durante la exportación.",
    )
    ws.cell(footer_row, 1).font = Font(color=INK_SOFT, italic=True, size=9)
    ws.cell(footer_row, 1).alignment = Alignment(horizontal="center", wrap_text=True)

    ws.print_area = f"A1:H{footer_row}"
    ws.print_title_rows = f"1:{item_header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    return wb
