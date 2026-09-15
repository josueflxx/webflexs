"""Formatted XLSX exports for internal commercial documents."""

from decimal import Decimal

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.services.sales_documents import build_internal_document_display_items


ORANGE = "FF6B35"
INK = "172033"
PAPER = "FFFFFF"
SURFACE = "F3F6FA"
LINE = "CBD5E1"
MONEY_FORMAT = '$ #,##0.00;[Red]-$ #,##0.00'
QUANTITY_FORMAT = "0.###"


def build_internal_export_filename(document):
    label = slugify(document.commercial_type_label or document.doc_type) or "documento"
    number = slugify(str(document.display_number)) or str(document.pk)
    return f"{label}-{number}.xlsx"


def build_internal_workbook(document):
    order = document.order
    client = document.client_profile or getattr(document.client_company_ref, "client_profile", None)
    rows = build_internal_document_display_items(document)
    wb = Workbook()
    ws = wb.active
    ws.title = "Documento"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A10"

    thin = Side(style="thin", color=LINE)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=INK)
    orange_fill = PatternFill("solid", fgColor=ORANGE)
    soft_fill = PatternFill("solid", fgColor=SURFACE)
    for col, width in {"A": 18, "B": 14, "C": 48, "D": 18, "E": 16, "F": 20}.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:F1")
    ws["A1"] = document.company.legal_name or document.company.name
    ws["A1"].font = Font(color=PAPER, bold=True, size=18)
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:F2")
    ws["A2"] = f"{document.commercial_type_label}  {document.display_number}"
    ws["A2"].font = Font(color=PAPER, bold=True, size=14)
    ws["A2"].fill = orange_fill
    ws["A2"].alignment = Alignment(vertical="center")

    issued_at = document.issued_at
    if issued_at and timezone.is_aware(issued_at):
        issued_at = timezone.localtime(issued_at).replace(tzinfo=None)
    info = [
        ("Cliente", getattr(client, "company_name", "") or getattr(order, "client_company", "") or "-"),
        ("CUIT / DNI", getattr(client, "cuit_dni", "") or getattr(order, "client_cuit", "") or "-"),
        ("Fecha", issued_at),
        ("Estado", "Anulado" if document.is_cancelled else "Emitido"),
    ]
    for index, (label, value) in enumerate(info, start=4):
        ws.cell(index, 1, label)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=6)
        ws.cell(index, 2, value)
        for col in range(1, 7):
            cell = ws.cell(index, col)
            cell.fill = soft_fill
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(index, 1).font = Font(color=INK, bold=True)

    headers = ["Código / SKU", "Cantidad", "Producto", "Precio unitario", "Bonificación", "Importe"]
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(9, col, label)
        cell.fill = header_fill
        cell.font = Font(color=PAPER, bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    current_row = 10
    for item in rows:
        if isinstance(item, dict):
            getter = item.get
        else:
            getter = lambda key, default=None, obj=item: getattr(obj, key, default)
        values = [
            getter("product_sku", "-") or "-",
            float(Decimal(str(getter("quantity", 0) or 0))),
            getter("product_name", "-") or "-",
            float(Decimal(str(getter("price_at_purchase", 0) or 0))),
            float(Decimal(str(getter("discount_percentage_used", 0) or 0)) / Decimal("100")),
            float(Decimal(str(getter("subtotal", 0) or 0))),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(current_row, col, value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col == 3)
            if col in {4, 6}:
                cell.number_format = MONEY_FORMAT
            elif col == 2:
                cell.number_format = QUANTITY_FORMAT
            elif col == 5:
                cell.number_format = "0.00%"
        current_row += 1

    if not rows:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        ws.cell(current_row, 1, "Este documento no tiene productos.")
        ws.cell(current_row, 1).alignment = Alignment(horizontal="center")
        current_row += 1

    totals = [
        ("Subtotal", getattr(order, "subtotal", Decimal("0.00"))),
        ("Descuento", getattr(order, "discount_amount", Decimal("0.00"))),
        ("TOTAL", getattr(order, "total", Decimal("0.00"))),
    ]
    current_row += 1
    for label, value in totals:
        ws.cell(current_row, 5, label)
        ws.cell(current_row, 6, float(Decimal(str(value or 0))))
        ws.cell(current_row, 5).font = Font(color=INK, bold=True)
        ws.cell(current_row, 6).font = Font(color=ORANGE if label == "TOTAL" else INK, bold=True)
        ws.cell(current_row, 6).number_format = MONEY_FORMAT
        current_row += 1

    ws.auto_filter.ref = f"A9:F{max(9, current_row - 5)}"
    return wb
