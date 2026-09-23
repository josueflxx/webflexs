"""Regression coverage for current admin prices across client-facing channels."""
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from tempfile import TemporaryDirectory

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import ClientCompany, ClientProfile
from catalog.models import Category, PriceList, PriceListItem, Product
from core.models import Company, CatalogExcelTemplate, CatalogExcelTemplateColumn, CatalogExcelTemplateSheet
from core.services.catalog_excel_exporter import build_catalog_workbook
from core.services.catalog_excel_status import latest_catalog_excel_source_change
from core.services.pricing import build_price_list_item_map, calculate_cart_pricing, get_base_price_for_product
from orders.models import Cart, CartItem
from catalog.tests_excel_utils import install_export_test_worker, download_generated_excel, excel_response_bytes


class CurrentCatalogPricingTests(TestCase):
    def setUp(self):
        install_export_test_worker(self)
        cache.clear()
        temporary_media = TemporaryDirectory()
        self.addCleanup(temporary_media.cleanup)
        override = self.settings(MEDIA_ROOT=temporary_media.name)
        override.enable()
        self.addCleanup(override.disable)
        self.company = Company.objects.create(name="Pricing regression", slug="pricing-regression")
        self.base_list = PriceList.objects.create(company=self.company, name="Base", slug="base")
        self.company.default_price_list = self.base_list
        self.company.save(update_fields=["default_price_list"])
        category = Category.objects.create(name="Pricing category", slug="pricing-category", visible_in_catalog=True)
        self.product = Product.objects.create(
            sku="BA04001", name="Buje regression", price=Decimal("7250.66"), category=category,
        )
        self.product.categories.add(category)
        self.old_item = PriceListItem.objects.create(
            price_list=self.base_list, product=self.product, price=Decimal("27762.62"),
        )
        self.user = User.objects.create_user(username="pricing-client", password="test-only")
        profile = ClientProfile.objects.create(user=self.user, company_name="Pricing client", is_approved=True)
        self.link = ClientCompany.objects.create(client_profile=profile, company=self.company)
        self.client.force_login(self.user)
        session = self.client.session
        session["active_company_id"] = self.company.pk
        session.save()
        self.template = CatalogExcelTemplate.objects.create(
            name="Pricing regression", is_active=True, is_client_download_enabled=True,
        )
        sheet = CatalogExcelTemplateSheet.objects.create(template=self.template, name="Pricing category")
        sheet.categories.add(category)
        for order, key in enumerate(("sku", "name", "price"), 1):
            CatalogExcelTemplateColumn.objects.create(sheet=sheet, key=key, order=order)

    def excel_price(self, response):
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(excel_response_bytes(response)), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                for row in sheet.iter_rows(values_only=True):
                    if row[0] == self.product.sku:
                        return Decimal(str(row[2]))
        finally:
            workbook.close()
        self.fail("Product was not found in the downloaded catalog")

    def assert_client_prices(self, base_price, final_price, net_base_price=None):
        if net_base_price is None:
            net_base_price = base_price
        listing = self.client.get(reverse("catalog"), {"q": self.product.sku})
        self.assertEqual(listing.status_code, 200)
        listed = next(p for p in listing.context["page_obj"] if p.pk == self.product.pk)
        self.assertEqual(listed.base_price, base_price)
        self.assertEqual(listed.final_price, final_price)
        detail = self.client.get(reverse("product_detail", args=[self.product.sku]))
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.context["base_price"], base_price)
        self.assertEqual(detail.context["final_price"], final_price)
        self.assertEqual(self.excel_price(download_generated_excel(self)), final_price)
        cart, _ = Cart.objects.get_or_create(user=self.user, company=self.company)
        CartItem.objects.get_or_create(cart=cart, product=self.product, defaults={"quantity": 1})
        cart_pricing = calculate_cart_pricing(cart, user=self.user, company=self.company)
        self.assertEqual(cart_pricing["items"][0].unit_price, net_base_price)
        self.assertEqual(cart_pricing["subtotal"], net_base_price)
        self.assertEqual(cart_pricing["discount_percentage"], self.link.discount_percentage)

    def test_base_list_follows_admin_across_web_excel_and_cart(self):
        self.assert_client_prices(Decimal("8773.30"), Decimal("8773.30"), net_base_price=Decimal("7250.66"))
        self.product.price = Decimal("8100.50")
        self.product.save(update_fields=["price", "updated_at"])
        self.assert_client_prices(Decimal("9801.61"), Decimal("9801.61"), net_base_price=Decimal("8100.50"))
        self.old_item.refresh_from_db()
        self.assertEqual(self.old_item.price, Decimal("27762.62"))

    def test_client_discount_is_preserved_with_matching_excel_rounding(self):
        self.link.discount_percentage = Decimal("25.00")
        self.link.save(update_fields=["discount_percentage"])
        self.assert_client_prices(Decimal("8773.30"), Decimal("6579.98"), net_base_price=Decimal("7250.66"))

    def test_custom_list_keeps_negotiated_prices(self):
        custom = PriceList.objects.create(company=self.company, name="Custom", slug="custom")
        PriceListItem.objects.create(price_list=custom, product=self.product, price=Decimal("6400.00"))
        self.link.price_list = custom
        self.link.discount_percentage = Decimal("15.00")
        self.link.save(update_fields=["price_list", "discount_percentage"])
        self.assert_client_prices(Decimal("7744.00"), Decimal("6582.40"), net_base_price=Decimal("6400.00"))

    def test_base_list_ignores_previously_loaded_snapshot_map(self):
        stale_map = {self.product.pk: self.old_item}
        self.assertEqual(get_base_price_for_product(self.product, self.base_list, stale_map), self.product.price)
        with self.assertNumQueries(0):
            self.assertEqual(build_price_list_item_map(self.base_list, [self.product.pk]), {})

    def test_base_list_zero_price_is_not_exported_from_a_positive_snapshot(self):
        self.product.price = Decimal("0.00")
        self.product.save(update_fields=["price", "updated_at"])
        workbook, stats = build_catalog_workbook(self.template, price_list=self.base_list)
        self.assertEqual(stats["total_rows"], 0)
        self.assertNotIn(self.product.sku, [cell.value for sheet in workbook for row in sheet for cell in row])

    def test_list_price_updates_invalidate_catalog_cache_timestamp(self):
        before = latest_catalog_excel_source_change(self.template)
        changed_at = max(before, timezone.now()) + timedelta(seconds=2)
        PriceListItem.objects.filter(pk=self.old_item.pk).update(price=Decimal("9999.00"), updated_at=changed_at)
        self.assertEqual(latest_catalog_excel_source_change(self.template), changed_at)
