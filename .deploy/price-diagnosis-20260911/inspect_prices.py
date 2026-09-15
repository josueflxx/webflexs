"""Read-only production pricing comparison; no client identities are printed."""
import hashlib
import json
import os
import sys
from pathlib import Path

sys.path.insert(0, "/var/www/webflexs")
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "flexs_project.settings.production")
import django
django.setup()

from django.db import connection, transaction
from django.db.models import Count, F, Max
from accounts.models import ClientCompany, ClientCategoryCompanyRule
from catalog.models import PriceList, PriceListItem, Product
from core.models import Company
from django.conf import settings

def emit(label, value):
    print(label + ": " + json.dumps(value, default=str, ensure_ascii=False))

with transaction.atomic():
    with connection.cursor() as cursor:
        cursor.execute("SET TRANSACTION READ ONLY")
    if "--url-product" in sys.argv:
        from django.urls import resolve
        from core.services.pricing import get_product_pricing
        from core.services.catalog_excel_exporter import _resolve_product_export_price
        from django.db.models import Q
        route = resolve("/admin-panel/productos/5957/")
        emit("url_route", {"view": route.view_name, "kwargs": route.kwargs})
        companies = list(Company.objects.select_related("default_price_list").order_by("id"))
        for product in Product.objects.filter(Q(pk=5957) | Q(sku__iexact="BA10001")).order_by("pk"):
            emit("compared_product", {key: getattr(product, key) for key in (
                "id", "sku", "name", "price", "cost", "iva_rate", "updated_at")})
            for company in companies:
                pricing = get_product_pricing(product, company=company)
                price_map = dict(PriceListItem.objects.filter(price_list=company.default_price_list, product=product).values_list("product_id", "price"))
                emit("compared_prices", {"id": product.pk, "sku": product.sku, "company": company.name,
                    "admin": product.price, "web_before_client_discount": pricing.base_price,
                    "excel_before_client_discount": _resolve_product_export_price(product, price_map=price_map)})
        template_path = Path("/var/www/webflexs/admin_panel/templates/admin_panel/products/workspace.html")
        if template_path.exists():
            lines = template_path.read_text().splitlines()
            emit("workspace_prices", [(i, line.strip()) for i, line in enumerate(lines, 1) if "product.price" in line or "final_price|" in line])
        sys.exit(0)
    if "--focus" in sys.argv:
        from core.services.pricing import get_product_pricing
        from core.services.catalog_excel_exporter import _resolve_product_export_price
        from openpyxl import load_workbook
        product = Product.objects.get(sku="BA10001")
        emit("reported_product_details", {
            key: getattr(product, key) for key in ("id", "sku", "name", "cost", "price", "iva_rate", "updated_at")
        })
        emit("similar_products", list(Product.objects.filter(name__icontains="DKW").values("sku", "name", "price", "cost")[:15]))
        for company in Company.objects.select_related("default_price_list").order_by("id"):
            pricing = get_product_pricing(product, company=company)
            price_map = dict(PriceListItem.objects.filter(price_list=company.default_price_list, product=product).values_list("product_id", "price"))
            emit("web_and_excel_resolution", {"company": company.name, "admin": product.price, "web": pricing.final_price,
                "discount": pricing.discount_percentage, "excel": _resolve_product_export_price(product, price_map=price_map)})
        for path in sorted((Path(settings.MEDIA_ROOT) / "catalog_exports").glob("catalogo_client_*.xlsx")):
            workbook = load_workbook(path, read_only=True, data_only=True)
            matches = []
            for sheet in workbook:
                if sheet.title == "INDICE":
                    continue
                for row in sheet.iter_rows(values_only=True):
                    if "BA10001" in row:
                        matches.append({"sheet": sheet.title, "row": row[:5]})
            emit("cached_product", {"file": path.name, "matches": matches})
            workbook.close()
        for relative_path in ("admin_panel/templates/admin_panel/products/list.html", "admin_panel/templates/admin_panel/products/form.html"):
            lines = (Path("/var/www/webflexs") / relative_path).read_text().splitlines()
            emit("admin_display", {"file": relative_path, "lines": [(number, line.strip()) for number, line in enumerate(lines, 1) if "product.price|" in line]})
        sys.exit(0)
    emit("companies", list(Company.objects.values("id", "name", "slug", "is_active", "default_price_list_id")))
    for price_list in PriceList.objects.order_by("id"):
        items = PriceListItem.objects.filter(price_list=price_list)
        mismatches = items.exclude(price=F("product__price"))
        visible_mismatches = mismatches.filter(product_id__in=Product.catalog_visible().values("pk"))
        emit("price_list", {
            "id": price_list.pk, "name": price_list.name, "company_id": price_list.company_id,
            "active": price_list.is_active, "item_count": items.count(),
            "different_from_admin": mismatches.count(),
            "visible_different": visible_mismatches.count(),
            "last_item_update": items.aggregate(value=Max("updated_at"))["value"],
            "samples": list(visible_mismatches.order_by("-product__updated_at").values(
                "product__sku", "product__name", "product__price", "price",
                "product__updated_at", "updated_at")[:8]),
        })
    emit("reported_product", list(Product.objects.filter(sku__iexact="BA10001").values("id", "sku", "name", "price", "updated_at")))
    emit("reported_product_list_prices", list(PriceListItem.objects.filter(product__sku__iexact="BA10001").values(
        "price_list_id", "price_list__name", "price", "created_at", "updated_at")))
    emit("client_pricing_counts", list(ClientCompany.objects.values(
        "company_id", "price_list_id", "discount_percentage"
    ).annotate(count=Count("id")).order_by("company_id", "price_list_id", "discount_percentage")))
    emit("category_rules", list(ClientCategoryCompanyRule.objects.values(
        "company_id", "client_category_id", "price_list_id", "discount_percentage", "is_active")))
    emit("latest_product_update", Product.objects.aggregate(value=Max("updated_at")))
    from django.conf import settings
    cache_path = Path(settings.MEDIA_ROOT) / "catalog_exports"
    emit("cached_catalogs", [
        {"file": path.name, "mtime": path.stat().st_mtime, "bytes": path.stat().st_size}
        for path in sorted(cache_path.glob("catalogo_client_*.xlsx"))
    ] if cache_path.exists() else [])
    for name in ("core/services/pricing.py", "core/services/catalog_excel_status.py", "core/services/catalog_excel_exporter.py"):
        path = Path("/var/www/webflexs") / name
        emit("source_hash", {"path": name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
