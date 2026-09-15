"""Verify published pricing using read-only DB access; optionally rebuild derived XLSX caches."""
import argparse
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import re
import sys
import time
from decimal import Decimal

APP = Path("/var/www/webflexs")
sys.path.insert(0, str(APP))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "flexs_project.settings.production")

parser = argparse.ArgumentParser()
parser.add_argument("mode", choices=["inspect", "verify", "warm"])
parser.add_argument("--staged", action="store_true")
args = parser.parse_args()

import django
django.setup()

if args.staged:
    for name in ("pricing", "catalog_excel_exporter", "catalog_excel_status"):
        module_name = f"core.services.{name}"
        spec = importlib.util.spec_from_file_location(module_name, Path(__file__).parent / "package" / "core" / "services" / f"{name}.py")
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)

from django.conf import settings
from django.db import connection, transaction
from accounts.models import ClientCompany, ClientCategoryCompanyRule
from catalog.models import PriceList, PriceListItem, Product
from core.models import CatalogExcelTemplate, Company
from core.services.pricing import build_price_list_item_map, calculate_final_price, get_base_price_for_product, get_product_pricing
from core.services.catalog_excel_exporter import build_catalog_workbook
from openpyxl import load_workbook


def fingerprint(queryset):
    digest = hashlib.sha256()
    for row in queryset.iterator(chunk_size=1000):
        digest.update((json.dumps(row, default=str, separators=(",", ":")) + "\n").encode())
    return digest.hexdigest()


def build_and_verify(template, price_list, discount):
    # Repeatable-read ensures the expected prices and workbook see the same data.
    expected = dict(Product.objects.values_list("sku", "price"))
    workbook, stats = build_catalog_workbook(template, price_list=price_list, discount_percentage=discount)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    saved = load_workbook(output, read_only=True, data_only=True)
    checked = 0
    samples = {}
    for sheet in saved:
        if sheet.title == "INDICE":
            continue
        for row in sheet.iter_rows(values_only=True):
            if len(row) < 3 or row[0] not in expected:
                continue
            target = calculate_final_price(expected[row[0]], discount)
            assert Decimal(str(row[2])) == target, (row[0], row[2], str(target))
            if row[0] in ("BA04001", "BA10001"):
                samples[row[0]] = str(target)
            checked += 1
    saved.close()
    assert checked > 6000, ("Unexpectedly small catalog", checked, stats)
    assert "BA04001" in samples, "Reported product missing"
    print(json.dumps({"check": "xlsx", "list": price_list.pk, "discount": str(discount), "verified_rows": checked, "samples": samples}), flush=True)
    return output.getvalue()


snapshot_started = time.time()
with transaction.atomic():
    with connection.cursor() as cursor:
        cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
    companies = list(Company.objects.filter(is_active=True).select_related("default_price_list").order_by("id"))
    print(json.dumps({"lists": list(PriceList.objects.values("id", "slug", "company_id", "name"))}), flush=True)
    if args.mode == "inspect":
        print(json.dumps({"fingerprints": {
            "products": fingerprint(Product.objects.order_by("pk").values_list("pk", "price")),
            "list_items": fingerprint(PriceListItem.objects.order_by("pk").values_list("pk", "price_list_id", "product_id", "price")),
            "client_rules": fingerprint(ClientCompany.objects.order_by("pk").values_list("pk", "company_id", "price_list_id", "discount_percentage")),
            "category_rules": fingerprint(ClientCategoryCompanyRule.objects.order_by("pk").values_list("pk", "price_list_id", "discount_percentage")),
        }}), flush=True)
    else:
        products = list(Product.objects.only("id", "sku", "price"))
        for company in companies:
            price_list = company.default_price_list
            assert price_list is not None and price_list.slug == "base", "Unexpected default list configuration"
            item_map = build_price_list_item_map(price_list, [p.pk for p in products])
            assert all(get_base_price_for_product(p, price_list, item_map) == p.price for p in products)
            samples = {}
            for product in products:
                if product.sku in ("BA04001", "BA10001"):
                    pricing = get_product_pricing(product, company=company)
                    assert pricing.base_price == product.price
                    samples[product.sku] = str(pricing.base_price)
            print(json.dumps({"check": "web_price_resolver", "company": company.name, "products_checked": len(products), "samples": samples}), flush=True)

        template = CatalogExcelTemplate.objects.filter(is_active=True, is_client_download_enabled=True).order_by("-updated_at", "id").first()
        assert template is not None
        if args.mode == "verify":
            build_and_verify(template, companies[0].default_price_list, Decimal("0"))
            build_and_verify(template, companies[0].default_price_list, Decimal("25"))
        else:
            cache_dir = Path(settings.MEDIA_ROOT) / "catalog_exports"
            cache_dir.mkdir(exist_ok=True)
            tasks = {}
            for path in cache_dir.glob("catalogo_client_*.xlsx"):
                match = re.fullmatch(r"catalogo_client_tpl(\d+)_plist(\d+)_disc(\d+)\.xlsx", path.name)
                if not match:
                    continue
                template_id, price_list_id, discount_hundredths = map(int, match.groups())
                if template_id == template.pk and PriceList.objects.filter(pk=price_list_id, slug="base").exists():
                    tasks[path.name] = (PriceList.objects.get(pk=price_list_id), Decimal(discount_hundredths) / 100)
            for company in companies:
                tasks.setdefault(f"catalogo_client_tpl{template.pk}_plist{company.default_price_list_id}_disc0.xlsx", (company.default_price_list, Decimal("0")))
            for name, (price_list, discount) in tasks.items():
                payload = build_and_verify(template, price_list, discount)
                target = cache_dir / name
                temporary = cache_dir / (name + ".pricing-fix.tmp")
                assert not temporary.exists(), temporary
                with temporary.open("xb") as file:
                    file.write(payload)
                os.chmod(temporary, 0o664)
                import grp
                import pwd
                os.chown(temporary, pwd.getpwnam("www-data").pw_uid, grp.getgrnam("www-data").gr_gid)
                os.utime(temporary, (snapshot_started, snapshot_started))
                os.replace(temporary, target)
                print(json.dumps({"cache_rebuilt": name, "bytes": len(payload)}), flush=True)
    print("PRICING_" + args.mode.upper() + "_OK", flush=True)
