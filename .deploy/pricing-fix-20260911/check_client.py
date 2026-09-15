"""Render the real client detail and serve a warmed download, without DB writes."""
import io
import json
import os
import sys
from decimal import Decimal
from unittest.mock import patch

sys.path.insert(0, "/var/www/webflexs")
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "flexs_project.settings.production")
import django
django.setup()

from django.db import connection, transaction
from django.shortcuts import render as real_render
from django.test import RequestFactory
from django.urls import reverse
from openpyxl import load_workbook
from accounts.models import ClientCompany
from catalog.models import Product
from catalog import views
from core.models import Company
from core.services.pricing import resolve_effective_discount_percentage, resolve_pricing_context

with transaction.atomic():
    with connection.cursor() as cursor:
        cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
    for company in Company.objects.filter(is_active=True).order_by("pk"):
        matched = None
        links = ClientCompany.objects.filter(
            company=company, is_active=True, discount_percentage=0, price_list__isnull=True,
            client_profile__is_approved=True, client_profile__user__is_active=True,
        ).select_related("client_profile__user").order_by("pk")
        for link in links:
            user = link.client_profile.user
            profile, client_company, category = resolve_pricing_context(user, company)
            discount = resolve_effective_discount_percentage(profile, company, client_company, category)
            if discount == 0:
                matched = user
                break
        assert matched is not None, "No zero-discount client available for check"
        for sku in ("BA04001", "BA10001"):
            product = Product.objects.get(sku=sku)
            request = RequestFactory().get(reverse("product_detail", args=[sku]), HTTP_HOST="flexsrepuestos.shop")
            request.user = matched
            request.session = {"active_company_id": company.pk}
            with patch.object(views, "render", wraps=real_render) as rendered:
                response = views.product_detail(request, sku)
            assert response.status_code == 200
            context = rendered.call_args.args[2]
            assert context["base_price"] == product.price and context["final_price"] == product.price
            print(json.dumps({"check": "rendered_client_detail", "company": company.name, "sku": sku, "price": str(context["final_price"]), "http": response.status_code}), flush=True)

        request = RequestFactory().get(reverse("catalog_client_excel_download"), HTTP_HOST="flexsrepuestos.shop")
        request.user = matched
        request.session = {"active_company_id": company.pk}
        response = views.client_catalog_excel_download(request)
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
        samples = {}
        expected = dict(Product.objects.filter(sku__in=("BA04001", "BA10001")).values_list("sku", "price"))
        for sheet in workbook:
            for row in sheet.iter_rows(values_only=True):
                if row[0] in expected:
                    assert Decimal(str(row[2])) == expected[row[0]]
                    samples[row[0]] = str(row[2])
        workbook.close()
        assert len(samples) == 2
        print(json.dumps({"check": "served_client_download", "company": company.name, "prices": samples}), flush=True)
print("CLIENT_PAGES_AND_DOWNLOADS_OK", flush=True)
