"""Production checks: read-only diagnostics plus explicitly queued derived exports."""
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import time
from decimal import Decimal
from unittest.mock import patch

APP = Path("/var/www/webflexs")
sys.path.insert(0, str(APP))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "flexs_project.settings.production")
import django
django.setup()
from django.db import connection, transaction
from django.test import RequestFactory, override_settings
from django.contrib.auth.models import User
from catalog import views
from catalog.models import Product
from core.models import CatalogExcelTemplate, Company
from core.services.pricing import calculate_final_price

mode = sys.argv[1]
def timed_tree(module):
    start = time.perf_counter()
    rows = module.get_cached_category_tree_rows()
    return [(r["category"].pk, r["depth"], r["full_path"], r["children_count"]) for r in rows], round(time.perf_counter()-start, 4)

if mode == "preflight":
    spec = importlib.util.spec_from_file_location("catalog.performance_candidate", Path(__file__).parent / "package/catalog/views.py")
    candidate = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(candidate)
    with override_settings(CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache", "LOCATION": "catalog-preflight"}}), transaction.atomic():
        with connection.cursor() as cursor:
            cursor.execute("SET TRANSACTION READ ONLY")
            cursor.execute("SET LOCAL statement_timeout = '10s'")
        before, old_time = timed_tree(views)
        after, new_time = timed_tree(candidate)
        again, warm_time = timed_tree(candidate)
        assert before == after == again, "Category tree changed"
        assert new_time < old_time / 2, "Category optimization did not improve latency"
        print(json.dumps({"identical_categories": len(after), "before_seconds": old_time, "after_seconds": new_time, "warm_seconds": warm_time}), flush=True)
    sys.exit(0)

from core.services.catalog_excel_jobs import ensure_export, export_spec, export_path, export_ready
from openpyxl import load_workbook
template = CatalogExcelTemplate.get_client_download_template()
assert template is not None
for company in Company.objects.filter(is_active=True).order_by("pk"):
    for discount in (Decimal("0"), Decimal("25")):
        spec = export_spec(template, company.default_price_list, discount)
        if mode == "queue":
            print(json.dumps({"company": company.pk, "discount": str(discount), **ensure_export(spec)}), flush=True)
        elif mode == "verify":
            assert export_ready(spec), f"Export not ready: company={company.pk}, discount={discount}"
            with transaction.atomic():
                with connection.cursor() as cursor:
                    cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
                expected = dict(Product.objects.values_list("sku", "price"))
                workbook = load_workbook(export_path(spec), read_only=True, data_only=True)
                checked = 0
                samples = {}
                for sheet in workbook:
                    if sheet.title == "INDICE":
                        continue
                    for row in sheet.iter_rows(values_only=True):
                        if len(row) < 3 or row[0] not in expected:
                            continue
                        actual = Decimal(str(row[2]))
                        wanted = calculate_final_price(expected[row[0]], discount)
                        assert actual == wanted, f"Wrong exported price: {row[0]}"
                        if row[0] in {"BA04001", "BA10001"}:
                            samples[row[0]] = str(actual)
                        checked += 1
                workbook.close()
                assert checked > 6000
                assert len(samples) == 2
                print(json.dumps({"company": company.pk, "discount": str(discount), "verified_rows": checked, "samples": samples}), flush=True)
if mode == "verify":
    with transaction.atomic():
        with connection.cursor() as cursor:
            cursor.execute("SET TRANSACTION READ ONLY")
        rows, elapsed = timed_tree(views)
        print(json.dumps({"published_category_rows": len(rows), "seconds": elapsed}), flush=True)
    print("CATALOG_PERFORMANCE_AND_EXPORTS_OK", flush=True)
