"""Serve the published Excel through actual approved-client authorization."""
import io
import json
import os
import sys
import time
from decimal import Decimal

sys.path.insert(0, "/var/www/webflexs")
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "flexs_project.settings.production")
import django
django.setup()
from django.db import connection, transaction
from django.test import RequestFactory
from openpyxl import load_workbook
from accounts.models import ClientCompany
from catalog.models import Product
from catalog.views import client_catalog_excel_download
from core.models import Company
from core.services.pricing import resolve_effective_discount_percentage, resolve_pricing_context

with transaction.atomic():
    with connection.cursor() as cursor:
        cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
    for company in Company.objects.filter(is_active=True).order_by("pk"):
        user = None
        for link in ClientCompany.objects.filter(company=company, is_active=True,
            client_profile__is_approved=True, client_profile__user__is_active=True
        ).select_related("client_profile__user").order_by("pk"):
            candidate = link.client_profile.user
            profile, company_link, category = resolve_pricing_context(candidate, company)
            discount = resolve_effective_discount_percentage(profile, company, company_link, category)
            if discount == 0:
                user = candidate
                break
        assert user is not None
        def request(params):
            req = RequestFactory().get("/catalogo/descargar-excel/", params, HTTP_HOST="flexsrepuestos.shop")
            req.user = user
            req.session = {"active_company_id": company.pk}
            return req
        started = time.perf_counter()
        status = client_catalog_excel_download(request({"status": "1"}))
        assert json.loads(status.content)["status"] == "ready"
        status_time = time.perf_counter() - started
        response = client_catalog_excel_download(request({}))
        assert response.status_code == 200 and response.streaming
        assert "no-store" in response["Cache-Control"]
        data = b"".join(response.streaming_content)
        # Response.close emits request_finished and would close this diagnostic's
        # shared read-only transaction; only close its file until checks finish.
        response.file_to_stream.close()
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        expected = dict(Product.objects.filter(sku__in=("BA04001", "BA10001")).values_list("sku", "price"))
        samples = {}
        for sheet in workbook:
            if sheet.title == "INDICE":
                continue
            for row in sheet.iter_rows(values_only=True):
                if len(row) >= 3 and row[0] in expected:
                    assert Decimal(str(row[2])) == expected[row[0]]
                    samples[row[0]] = str(row[2])
        workbook.close()
        assert len(samples) == 2
        print(json.dumps({"company": company.pk, "authenticated_status_seconds": round(status_time, 4),
            "download_bytes": len(data), "verified_prices": samples}), flush=True)
print("AUTHENTICATED_CLIENT_DOWNLOADS_OK", flush=True)
