import io
import os
import sys

sys.path.insert(0, "/var/www/webflexs")

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "flexs_project.settings.production")

import django

django.setup()

from openpyxl import load_workbook

from core.models import CatalogExcelTemplate
from core.services.catalog_excel_exporter import build_catalog_workbook


template = (
    CatalogExcelTemplate.objects.filter(is_client_download_enabled=True)
    .order_by("id")
    .first()
)
if template is None:
    raise RuntimeError("No hay una plantilla de catalogo habilitada para clientes.")

workbook, stats = build_catalog_workbook(template)
if "ACERO" not in workbook.sheetnames:
    raise RuntimeError(f"La plantilla publica no genero la hoja ACERO: {workbook.sheetnames}")

worksheet = workbook["ACERO"]
if len(worksheet._images) != 1:
    raise RuntimeError(f"La hoja ACERO contiene {len(worksheet._images)} imagenes; se esperaba 1.")

output = io.BytesIO()
workbook.save(output)
output.seek(0)
saved_workbook = load_workbook(output)
saved_image = saved_workbook["ACERO"]._images[0]
if saved_image.anchor._from.row != 0 or saved_image.anchor._from.col < 4:
    raise RuntimeError("La guia de medidas no quedo ubicada a la derecha de la tabla.")

print(
    "CATALOG_MEASURE_GUIDE_OK",
    f"template={template.pk}",
    f"rows={stats.get('total_rows', 0)}",
    f"anchor_col={saved_image.anchor._from.col + 1}",
)
